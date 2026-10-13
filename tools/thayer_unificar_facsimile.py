#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Catálogo unificado de recortes do fac-símile de Thayer.
NÃO CONVERTE candidatos OCR em artigos visualmente atestados sem revisão.
"""
from __future__ import annotations
from collections import Counter,defaultdict
from datetime import datetime,timezone
import json,os,pathlib,subprocess,time,urllib.request,urllib.error
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter

ROOT=pathlib.Path(__file__).resolve().parents[1]
OUT=ROOT/"entregas/THAYER_catalogo_fotografico_UNIFICADO_provisorio.xlsx"
JOUT=ROOT/"entregas/THAYER_catalogo_fotografico_UNIFICADO_provisorio.json"
ROOT_URL="https://raw.githubusercontent.com/linguasbiblicasoriginais-rgb/scriptura-lexicon/"
PRI=ROOT_URL+"chat-gpt-thayer-indice-titulos/entregas/THAYER_inventario_fotografico_candidatos.json"
SEC=ROOT_URL+"chat-gpt-thayer-facsimile-apoio/entregas/THAYER_cabecalhos_recuperados_coluna_direita.json"
def load(url):
    req=urllib.request.Request(url+"?v="+str(int(time.time()//60)),headers={"User-Agent":"ScripturaLexicon-merged-scan-index/1.0","Cache-Control":"no-cache"})
    try:
        with urllib.request.urlopen(req,timeout=90) as f:return json.load(f)
    except (urllib.error.HTTPError,urllib.error.URLError,TimeoutError,json.JSONDecodeError):
        return None
def k(r):return (int(r["pagina_pdf"]),int(r["x_scan"]),int(r["y_scan"]))
def greek_count(source):
    return sum(bool(r.get("titulo_grego_ocr_fotografia")) for r in source.get("candidatos",[]))
def gather():
    # Esperar a gravação REAL de OCR e recuperação na origem.
    last_primary=last_supplement=None
    ready=False
    for i in range(44):
        last_primary=load(PRI)
        last_supplement=load(SEC)
        oldlen=len(last_primary.get("candidatos",[])) if last_primary else 0
        oldocr=greek_count(last_primary) if last_primary else 0
        supp=len(last_supplement.get("entries",[])) if last_supplement else 0
        print(f"VERIFICACAO {i}: originais={oldlen} OCR grego={oldocr} recuperados={supp}",flush=True)
        if oldlen>3000 and oldocr>1000 and supp>50:
            ready=True
            break
        if i<43:time.sleep(60)
    if not last_primary or not last_supplement:
        raise RuntimeError("Uma das bases geradas do fac-símile ainda não existe. Não publicar falsa completude.")
    return last_primary,last_supplement,ready
def write(primary,complement,ready):
    prim=primary["candidatos"]; suppl=complement["entries"]
    uniq={k(r):dict(r,origem="Passagem geométrica inicial") for r in prim}
    for r in suppl:
        key=k(r)
        if key not in uniq:
            uniq[key]=dict(r,origem="Correção da fronteira de coluna direita")
        elif not uniq[key].get("titulo_grego_ocr_fotografia") and r.get("titulo_grego_ocr_fotografia"):
            uniq[key]=dict(r,origem="Correção fotográfica (substituição de OCR vazio)")
    rows=list(uniq.values())
    rows.sort(key=lambda z:(z["pagina_pdf"],0 if z["coluna"]=="esquerda" else 1,z["y_scan"],z["x_scan"]))
    counts=Counter(r["pagina_impressa"] for r in rows)
    readcounts=Counter(r["pagina_impressa"] for r in rows if r.get("titulo_grego_ocr_fotografia"))
    n_ocr=sum(readcounts.values());miss=len(rows)-n_ocr
    w=Workbook()
    info=w.active;info.title="Estado e critérios"
    summaries=[
        ["THAYER — CATÁLOGO UNIFICADO BASEADO NAS FOTOGRAFIAS"],
        ["Fonte fotográfica","thayersgreekengl0000thay, edição impressa de Joseph Henry Thayer (digitalizada pelo Internet Archive)"],
        ["Folhas do corpo lexicográfico","PDF 25–707; 683 páginas impressas"],
        ["Candidatos de cabeçalhos detectados",len(rows)],
        ["Candidatos com leitura grega OCR",n_ocr],
        ["Candidatos sem leitura grega OCR",miss],
        ["Páginas com 0 candidatos detectados",683-len(counts)],
        ["Passagem inicial",len(prim)],
        ["Recuperação adicional",len(suppl)],
        ["Combinação das duas passagens","SIM — chave geométrica folha + coordenadas; duplicações posicionais suprimidas"],
        ["ORIGEM DA GRAFIA","OCR grego feito a partir de recortes do fac-símile; transcrição automática com possibilidade de erros"],
        ["ARTIGOS VISUALMENTE CONFERIDOS","NÃO CERTIFICADOS — nenhum candidato automático foi homologado manualmente pelo script"],
        ["QUANTIDADE DEFINITIVA DOS VERBETES","INDETERMINADA — candidatos não equivalem a artigos: falsos positivos e omissões remanescentes são possíveis"],
        ["Marcadores Strong","Não usados para gerar os títulos"],
        ["5.594 vocábulos segundo Thayer","É o vocabulário do NT declarado no prefácio; não é automaticamente número de artigos impressos"],
        ["Fonte de coordenadas","DjVu XML da mesma digitalização do fac-símile"],
        ["Situação dos trabalhos","Duas passagens OCR recebidas" if ready else "PROCESSAMENTO INCOMPLETO — uma das duas passagens não terminou"],
        ["Data da edição",datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")]
    ]
    for line in summaries:info.append(line)
    master=w.create_sheet("Títulos candidatos — OCR grego")
    headers=["Nº","Lema grego lido no fac-símile por OCR","Página impressa","Folha PDF","Coluna",
        "Posição X","Posição Y","Confiança OCR","Situação","Trecho OCR de base","Origem geométrica","Fotografia"]
    master.append(headers)
    dubious=w.create_sheet("Leituras ausentes ou incertas")
    dubious.append(headers)
    for i,r in enumerate(rows,1):
        link=f'https://archive.org/details/thayersgreekengl0000thay/page/n{r["pagina_pdf"]-1}/mode/1up'
        v=[i,r.get("titulo_grego_ocr_fotografia") or "",r["pagina_impressa"],r["pagina_pdf"],
           r["coluna"],r["x_scan"],r["y_scan"],r.get("confianca_ocr_grego") or "",
           r["status"],r["texto_bruto_inicio"],r["origem"],"Abrir fotografia"]
        master.append(v)
        c=master.cell(master.max_row,12);c.hyperlink=link;c.font=Font(color="14648D",underline="single")
        conf=r.get("confianca_ocr_grego")
        try:low=conf is None or float(conf)<60
        except (ValueError,TypeError):low=True
        if low or not r.get("titulo_grego_ocr_fotografia"):
            dubious.append(v)
            dubious.cell(dubious.max_row,12).hyperlink=link
    coverage=w.create_sheet("Cobertura — 683 páginas")
    coverage.append(["Página impressa","Folha PDF","Candidatos (sem revisão)","Com leitura OCR","Pendentes de leitura OCR",
        "Conferência visual integral","Fotografia"])
    for p in range(1,684):
        coverage.append([p,p+24,counts[p],readcounts[p],counts[p]-readcounts[p],
                         "NÃO CONCLUÍDA",f'https://archive.org/details/thayersgreekengl0000thay/page/n{p+23}/mode/1up'])
        url=coverage.cell(coverage.max_row,7)
        url.hyperlink=url.value;url.font=Font(color="14648D",underline="single")
    for sh in [master,dubious,coverage]:
        sh.freeze_panes="C2";sh.sheet_view.showGridLines=False;sh.auto_filter.ref=sh.dimensions
        sh.row_dimensions[1].height=36
        for c in sh[1]:
            c.fill=PatternFill("solid",fgColor="163B52")
            c.font=Font(color="FFFFFF",bold=True)
            c.alignment=Alignment(wrap_text=True,vertical="center")
    for s in [master,dubious]:
        for col,v in enumerate([9,44,21,17,15,14,14,19,52,75,45,23],1):
            s.column_dimensions[get_column_letter(col)].width=v
    for i,width in enumerate([20,18,27,24,26,36,85],1):
        coverage.column_dimensions[get_column_letter(i)].width=width
    info.column_dimensions["A"].width=47;info.column_dimensions["B"].width=115
    info.sheet_view.showGridLines=False
    for cells in info:
        cells[0].font=Font(bold=True,color="163B52")
        if len(cells)>1:cells[1].alignment=Alignment(wrap_text=True)
    w.properties.title="THAYER: índice provisório extraído das fotografias — duas passagens"
    w.properties.creator="Scriptura Lexicon"
    OUT.parent.mkdir(parents=True,exist_ok=True)
    w.save(OUT)
    JOUT.write_text(json.dumps({"source":"thayersgreekengl0000thay","folios_scanned":683,"candidates":len(rows),
      "greek_OCR_detected":n_ocr,"unread":miss,"manual_review_completed":False,
      "full_machine_passes_finished":ready,"entries":rows},ensure_ascii=False)+"\n",encoding="utf-8")
    print(f"GERADO: {OUT.name} | candidatos={len(rows)} | com_OCR={n_ocr} | sem_OCR={miss}",flush=True)
def publish():
    subprocess.run(["git","config","user.name","github-actions[bot]"],check=True,cwd=ROOT)
    subprocess.run(["git","config","user.email","41898282+github-actions[bot]@users.noreply.github.com"],check=True,cwd=ROOT)
    subprocess.run(["git","add",str(OUT.relative_to(ROOT)),str(JOUT.relative_to(ROOT))],check=True,cwd=ROOT)
    subprocess.run(["git","commit","-m","THAYER: consolidar índices OCR de 683 folhas do fac-símile (provisório)"],check=True,cwd=ROOT)
    subprocess.run(["git","push","origin","HEAD:chat-gpt-thayer-catalogo-fotografico"],check=True,cwd=ROOT)
if __name__=="__main__":
    main1,main2,ready=gather()
    write(main1,main2,ready)
    publish()
