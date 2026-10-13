#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Inventario de cabecalhos do fac-simile Thayer, SEM substituir por Strong ou dataset auxiliar.
O OCR da imagem produz candidatos, e nao a certificacao de que foram conferidos por humano.
"""
import csv
import datetime
import html
import io
import math
import os
from pathlib import Path
import re
import subprocess
import sys
import time
import unicodedata
import urllib.request

import fitz
import pandas as pd
import pytesseract
from PIL import Image, ImageOps, ImageEnhance
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.worksheet.table import Table,TableStyleInfo

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/"entregas/THAYER_titulos_facsimile_OCR_provisorio.xlsx"
PDF=Path("/tmp/thayersgreekengl0000thay.pdf")
ARCHIVE="https://archive.org/download/thayersgreekengl0000thay/thayersgreekengl0000thay.pdf"
PDF_START=25
PDF_LAST=707
WORD_RE=re.compile(r"[\u0370-\u03ff\u1f00-\u1fff]")
GREEK=re.compile(r"^[\u0370-\u03ff\u1f00-\u1fff\u0300-\u036f'’᾿\u002d]+")
def getpdf():
    if PDF.exists() and PDF.stat().st_size>60_000_000:return
    print("Baixando digitalizacao do fac-simile de identificador thayersgreekengl0000thay",flush=True)
    req=urllib.request.Request(ARCHIVE,headers={"User-Agent":"ScripturaLexicon-facsimile-catalogue/1.0"})
    with urllib.request.urlopen(req,timeout=180) as src,open(PDF,"wb") as dst:
        while True:
            block=src.read(1024*1024)
            if not block:break
            dst.write(block)
    print("Tamanho digitalizacao",PDF.stat().st_size,flush=True)
def clean(s):
    return re.sub(r"\s+"," ",str(s)).strip()
def img_page(page):
    # Duas colunas. Usar amostragem suficiente para conferir cabecalhos pequenos.
    pix=page.get_pixmap(matrix=fitz.Matrix(2.45,2.45),colorspace=fitz.csGRAY,alpha=False)
    return Image.frombytes("L",[pix.width,pix.height],pix.samples)
def crop_columns(im):
    w,h=im.size
    # margens do exemplar (corpo do texto vai de aproximadamente 6% a 93% da largura)
    top=int(h*.050)
    bottom=int(h*.955)
    cols=[("esquerda",int(w*.065),int(w*.500)),("direita",int(w*.505),int(w*.945))]
    for side,x1,x2 in cols:
        region=im.crop((x1,top,x2,bottom))
        region=ImageEnhance.Contrast(region).enhance(1.55)
        yield side,x1,top,region
def read_candidates(pdf_page):
    doc=fitz.open(PDF)
    page=doc[pdf_page-1]
    # PDF texto, quando presente, fica preservado apenas como testemunho bruto.
    ocr_reference=clean(page.get_text("text")[:210])
    im=img_page(page)
    res=[]
    for side,x_offset,y_offset,crop in crop_columns(im):
        d=pytesseract.image_to_data(crop,lang="grc+eng",config="--psm 4",output_type=pytesseract.Output.DICT)
        lines={}
        for i,word in enumerate(d["text"]):
            word=clean(word)
            if not word:continue
            key=(d["block_num"][i],d["par_num"][i],d["line_num"][i])
            lines.setdefault(key,[]).append(i)
        seen=set()
        for idxs in lines.values():
            idxs.sort(key=lambda i:d["left"][i])
            start=idxs[0]
            x=d["left"][start]
            # Evitar palavras no meio da linha. Capturar inicio de artigo / paragrafo.
            if x>int(crop.width*.13):continue
            all_words=[clean(d["text"][i]) for i in idxs]
            raw=" ".join(all_words)
            head=clean(all_words[0])
            if head and head[0] in "*†'‘’\"“([":
                head=head.lstrip("*†'‘’\"“([")
            leading=GREEK.match(head)
            if not leading:continue
            head=leading.group(0).strip("-'")
            if not WORD_RE.search(head) or len(head)<2 or len(head)>36:continue
            ratio=sum(bool(WORD_RE.match(ch)) for ch in head)/max(1,len(head))
            if ratio<.58:continue
            conf=float(d["conf"][start])
            if conf<0:conf=0
            y=d["top"][start]
            if y<0 or y>crop.height*.96:continue
            uniq=(side,round(y/7),head)
            if uniq in seen:continue
            seen.add(uniq)
            # Heuristica de cabecalho, transparencia: nem todo candidato sera um verbete.
            ctx=raw[:130]
            result={
                "pagina_pdf":pdf_page,"pagina_impressa":pdf_page-24,
                "coluna":side,"y_imagem":int(y+y_offset),"x_imagem":int(x+x_offset),
                "leitura_ocr":head,"contexto_da_linha":ctx,
                "confianca_ocr":round(conf,1),
                "classificacao":"CANDIDATO DE CABEÇALHO — NÃO CONFERIDO VISUALMENTE",
                "nota":"A palavra grega foi reconhecida no início de linha da imagem; a distinção entre cabeçalho e citação exige inspeção humana.",
                "imagem_pix_width":im.width,"imagem_pix_height":im.height,
            }
            # A pontuacao historica sugere cabecalho, mas NAO confirma.
            if re.search(r"[,;:]|[—-]",raw[:len(all_words[0])+12]):
                result["nota"]+=" Pontuação inicial compatível."
            res.append(result)
    return res,ocr_reference
def main():
    getpdf()
    doc=fitz.open(PDF)
    if len(doc)!=762:raise RuntimeError(f"A fonte baixada tem {len(doc)} paginas; esperadas 762. Abortando para evitar livro diverso.")
    text=clean(doc[24].get_text("text")).lower()
    if "testament" not in text and "lexicon" not in text:
        raise RuntimeError("A pagina PDF 25 nao tem evidencia da abertura NEW TESTAMENT LEXICON; fonte inconclusiva")
    doc.close()
    entries=[];pgs=[]
    from concurrent.futures import ThreadPoolExecutor,as_completed
    # Uma tarefa por pagina (imagem lida independentemente).
    with ThreadPoolExecutor(max_workers=3) as pool:
        tasks={pool.submit(read_candidates,p):p for p in range(PDF_START,PDF_LAST+1)}
        for i,future in enumerate(as_completed(tasks),1):
            page=tasks[future]
            try:
                candidates,txt=future.result()
                entries.extend(candidates)
                pgs.append([page,page-24,len(candidates),"OCR da imagem executado","",txt])
            except Exception as e:
                pgs.append([page,page-24,0,"FALHA DE PROCESSAMENTO",str(e)[:180],""])
            if i%50==0:print(f"Facsimile: {i}/{PDF_LAST-PDF_START+1} folhas processadas; {len(entries)} candidatos",flush=True)
    entries.sort(key=lambda z:(z["pagina_pdf"],0 if z["coluna"]=="esquerda" else 1,z["y_imagem"]))
    if len(entries)<100:
        raise RuntimeError(f"OCR visual encontrou apenas {len(entries)} candidatos, muito poucos; não publicar arquivo que simule inventario.")
    w=Workbook()
    main=w.active;main.title="Candidatos do fac-símile"
    headers=["Nº","Leitura grega OCR","Página impressa","Folha do PDF","Coluna","X imagem (px)","Y imagem (px)","Confiança OCR","Trecho da imagem reconhecido","Estado da leitura","Observações"]
    main.append(headers)
    for i,r in enumerate(entries,1):
        main.append([i,r["leitura_ocr"],r["pagina_impressa"],r["pagina_pdf"],r["coluna"],
            r["x_imagem"],r["y_imagem"],r["confianca_ocr"],r["contexto_da_linha"],r["classificacao"],r["nota"]])
    pc=w.create_sheet("Folhas do fac-símile")
    pc.append(["Folha do PDF","Página impressa","Inícios de linha gregos detectados","Processamento","Impedimento","Amostra da camada textual do PDF"])
    for row in sorted(pgs,key=lambda r:r[0]):pc.append(row)
    info=w.create_sheet("Método e ressalvas",0)
    facts=[
        ["THAYER | CATÁLOGO DE CANDIDATOS EXTRAÍDOS DAS IMAGENS"],
        ["Obra","Joseph Henry Thayer, A Greek-English Lexicon of the New Testament — Corrected Edition, American Book Company"],
        ["Fonte visual","Fac-símile do Internet Archive, identificador thayersgreekengl0000thay (mesma digitalização identificada no PDF enviado)"],
        ["Folhas digitais",762],["Primeira folha do texto lexical",25],["Última folha do texto lexical",707],
        ["Páginas lexicais processadas",len(pgs)],["Candidatos de cabeçalhos identificados",len(entries)],
        ["Número exato de verbetes da obra","NÃO DETERMINADO"],
        ["Método","As páginas fotografadas foram renderizadas em escala de cinza, separadas em colunas e submetidas a reconhecimento de grego+inglês. Foram retidos inícios de linha contendo caracteres gregos."],
        ["Limitação fundamental","São CANDIDATOS automáticos extraídos das fotografias; incluem potencialmente falsos positivos e omissões. Não foram cotejados manualmente em todas as 683 folhas."],
        ["Sem substituição silenciosa","NENHUM título foi extraído da base de Strong ou do dataset digital auxiliar para preencher lacunas."],
        ["Coordenadas","X/Y indicam a posição na página rasterizada com escala 2,45. Páginas seguem deslocamento: folha do PDF = página impressa + 24."],
        ["Como conferir","Abrir a folha do PDF indicada, localizar coluna e Y, comparar OCR ao original. Depois substituir por leitura humana e alterar status."],
        ["Exclusões","Prefácio, bibliografia, apêndices, folha sem texto principal."],
        ["Total declarado pelo autor","5.594 palavras no vocabulário do NT (p. XVIII da introdução) — não confundir com quantidade de artigos."],
        ["Data da extração",datetime.date.today().isoformat()]
    ]
    for row in facts:info.append(row)
    info.column_dimensions["A"].width=40;info.column_dimensions["B"].width=120
    for row in info:
        row[0].font=Font(color="164961",bold=True,size=11)
        if len(row)>1:row[1].alignment=Alignment(wrap_text=True,vertical="top")
    info.freeze_panes="A2"; info.sheet_view.showGridLines=False
    for ws,widths in [(main,[9,35,20,15,14,17,17,18,65,64,120]),(pc,[18,22,30,35,50,110])]:
        ws.freeze_panes="C2";ws.sheet_view.showGridLines=False
        for i,width in enumerate(widths,1):
            ws.column_dimensions[__import__("openpyxl").utils.get_column_letter(i)].width=width
        for c in ws[1]:
            c.fill=PatternFill("solid",fgColor="163D57")
            c.font=Font(color="FFFFFF",bold=True,size=10)
            c.alignment=Alignment(vertical="center",wrap_text=True)
        ws.row_dimensions[1].height=32
        ws.auto_filter.ref=ws.dimensions
    for row in main.iter_rows(min_row=2):
        row[1].font=Font(name="Palatino Linotype",size=12,color="153B55")
        row[9].font=Font(color="AD671E",size=10)
    main.sheet_view.zoomScale=90
    w.properties.title="THAYER: candidatos lexicais OCR do facsimile — não homologado"
    w.properties.subject="Extração preliminar visual automatizada, sem conferência filológica manual integral"
    w.properties.creator="Scriptura Lexicon"
    OUT.parent.mkdir(parents=True,exist_ok=True)
    w.save(OUT)
    subprocess.run(["git","config","user.name","github-actions[bot]"],cwd=ROOT,check=True)
    subprocess.run(["git","config","user.email","41898282+github-actions[bot]@users.noreply.github.com"],cwd=ROOT,check=True)
    subprocess.run(["git","add",str(OUT.relative_to(ROOT))],cwd=ROOT,check=True)
    subprocess.run(["git","commit","-m","THAYER: planilha de candidatos extraídos das fotografias do fac-símile"],cwd=ROOT,check=True)
    subprocess.run(["git","push","origin","HEAD:chat-gpt-thayer-indice-titulos"],cwd=ROOT,check=True)
    print(f"SUCESSO {len(entries)} candidatos de 683 páginas; {OUT.relative_to(ROOT)}",flush=True)
if __name__=="__main__":main()
