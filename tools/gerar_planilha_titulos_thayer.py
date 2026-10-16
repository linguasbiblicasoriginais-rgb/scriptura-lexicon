#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Exportar títulos transcritos de Thayer em XLSX de inventário.
Este roteiro NÃO afirma ter contado todos os cabeçalhos no fac-símile.
As leituras são transcrição auxiliar CC0; cotejo fotográfico pendente.
"""
from __future__ import annotations
import collections
import datetime as dt
import hashlib
import json
import pathlib
import re
import urllib.request

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

URL = "https://raw.githubusercontent.com/nigelmsipa/thayers-greek-lexicon-dataset/master/data/thayer_lexicon.jsonl"
PDF_REFERENCIA = "thayersgreekengl0000thay.pdf — Internet Archive, edição corrigida"
PDF_URL = "https://archive.org/details/thayersgreekengl0000thay"
DATASET_URL = "https://github.com/nigelmsipa/thayers-greek-lexicon-dataset"
DEST = pathlib.Path(__file__).resolve().parents[1] / "entregas" / "THAYER_indice_titulos_inventario.xlsx"
COR_PRIMARY = "143B50"
COR_SECONDARY = "286276"
COR_PALE = "E8F1F4"
COR_WARNING = "FFF2D8"
COR_TEXT = "213442"

def page_str(v):
    if v is None:
        return ""
    if isinstance(v,list):
        return "; ".join(str(item).strip() for item in v if str(item).strip())
    return str(v).strip()

def page_number(print_pages):
    if isinstance(print_pages,list):
        for x in print_pages:
            s = str(x).strip()
            m = re.match(r"^\d{1,4}$",s)
            if m:
                return int(s)
    return None

def fill_summary(ws, data, fingerprint):
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:G2")
    ws["A1"] = "THAYER | Índice dos títulos lexicográficos"
    ws["A1"].font = Font(name="Aptos Display", size=20, color="FFFFFF", bold=True)
    ws["A1"].fill = PatternFill("solid", fgColor=COR_PRIMARY)
    ws["A1"].alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 26
    ws.row_dimensions[2].height = 17
    labels=[
        (4, "Registros na transcrição auxiliar",len(data)),
        (5, "Títulos gregos distintos (contagem textual)", len({str(x.get("lemma") or "").strip() for x in data if str(x.get("lemma") or "").strip()})),
        (6, "Títulos sem leitura grega",sum(1 for x in data if not str(x.get("lemma") or "").strip())),
        (7, "Linhas com indicação de página impressa",sum(1 for x in data if page_number((x.get("source") or {}).get("print_pages")) is not None)),
        (8, "Total de verbetes do FAC-SÍMILE", "NÃO DETERMINADO"),
        (9, "Palavras do vocabulário segundo Thayer (p. XVIII)",5594),
        (10,"Palavras marcadas com asterisco segundo Thayer",5300),
    ]
    for row, label, value in labels:
        ws.merge_cells(start_row=row,start_column=1,end_row=row,end_column=4)
        c=ws.cell(row,1,label)
        c.font=Font(name="Aptos",color=COR_TEXT,bold=True,size=11)
        c.fill=PatternFill("solid",fgColor=COR_PALE)
        d=ws.cell(row,5,value)
        d.font=Font(name="Aptos",color=COR_PRIMARY,bold=True,size=12)
        d.alignment=Alignment(horizontal="right")
        d.number_format='#,##0' if isinstance(value,int) else '@'
    ws.merge_cells("A12:G14")
    ws["A12"]="ATENÇÃO — Este arquivo enumera os títulos da transcrição digital auxiliar (6.134 registros na versão consultada), não uma extração visual exaustiva dos cabeçalhos do fac-símile. Não atribua este total ao livro impresso sem cotejo integral."
    ws["A12"].alignment=Alignment(vertical="center",wrap_text=True,indent=1)
    ws["A12"].fill=PatternFill("solid",fgColor=COR_WARNING)
    ws["A12"].font=Font(name="Aptos",size=11,color="633F16",bold=True)
    ws.row_dimensions[12].height=36
    ws.merge_cells("A16:G18")
    ws["A16"]="Ordem das linhas: ordem da transcrição auxiliar, não reordenada. Lemas preservados conforme a fonte; nenhuma correção ortográfica ou normalização silenciosa. Homógrafos e duplicações permanecem separados."
    ws["A16"].alignment=Alignment(vertical="center",wrap_text=True,indent=1)
    ws["A16"].font=Font(name="Aptos",size=11,color=COR_TEXT)
    ws.merge_cells("A20:G22")
    ws["A20"]="Paginação: na cópia de 762 páginas fornecida, a página impressa 1 aparece na folha PDF 25. A indicação «folha estimada» soma 24 ao número da página impressa indicada na transcrição; não constitui cotejo página a página."
    ws["A20"].alignment=Alignment(vertical="center",wrap_text=True,indent=1)
    ws["A20"].font=Font(name="Aptos",size=11,color=COR_TEXT)
    extra=[
        ("Fac-símile informado",PDF_REFERENCIA),
        ("Digitalização pública correlata",PDF_URL),
        ("Transcrição digital auxiliar",DATASET_URL),
        ("Arquivo bruto transcrito",URL),
        ("SHA256 da fonte digital",fingerprint),
        ("Data da exportação",dt.date.today().isoformat()),
        ("Limitação", "OCR grego do PDF apresenta falhas; títulos NÃO foram cotejados integralmente nas fotografias"),
    ]
    for i,(k,v) in enumerate(extra,25):
        ws.cell(i,1,k).font=Font(name="Aptos",bold=True,color=COR_SECONDARY,size=10)
        ws.merge_cells(start_row=i,start_column=2,end_row=i,end_column=7)
        c=ws.cell(i,2,v)
        c.font=Font(name="Aptos",size=10,color=COR_TEXT)
        if isinstance(v,str) and v.startswith("https://"):
            c.hyperlink=v
            c.font=Font(name="Aptos",size=10,color="1E6A94",underline="single")
        c.alignment=Alignment(wrap_text=True,vertical="top")
    for col in "ABCDEFG": ws.column_dimensions[col].width=19
    ws.column_dimensions["A"].width=30
    ws.column_dimensions["B"].width=28
    ws.freeze_panes="A4"

def fill_index(ws,entries):
    ws.sheet_view.showGridLines=False
    headers=("Nº","Título (grego; transcrição auxiliar)","Identificador na fonte","Strong vinculado","Páginas impressas indicadas","Folhas PDF na fonte auxiliar","Folha no PDF enviado (estimada)","Conferência visual","Observações de extração")
    ws.append(list(headers))
    invalid=0
    for i,e in enumerate(entries,1):
        src=e.get("source") or {}
        lemma=str(e.get("lemma") or "").strip()
        if not lemma:invalid+=1
        p=page_number(src.get("print_pages") or [])
        estimated=p+24 if p is not None and p+24<=762 else None
        strong=e.get("strongs") or []
        if isinstance(strong,str):strong=[strong]
        notes=[]
        if not lemma:notes.append("Título ausente no registro digital")
        if not strong:notes.append("Sem índice Strong vinculado")
        if p is None:notes.append("Sem paginação impressa numérica")
        if lemma and not re.search(r"[\u0370-\u03ff\u1f00-\u1fff]",lemma):notes.append("Sem caráter grego identificável")
        ws.append([
            i,lemma,str(e.get("id") or ""),
            "; ".join(str(s) for s in strong),
            page_str(src.get("print_pages")),
            page_str(src.get("pdf_pages")),
            estimated,
            "PENDENTE — cotejo com imagem",
            "; ".join(notes),
        ])
    ws.freeze_panes="C2"
    ws.auto_filter.ref=f"A1:I{len(entries)+1}"
    ws.sheet_view.zoomScale=90
    ws.row_dimensions[1].height=32
    widths={"A":10,"B":38,"C":23,"D":25,"E":28,"F":30,"G":32,"H":33,"I":55}
    for col,w in widths.items():ws.column_dimensions[col].width=w
    for c in ws[1]:
        c.font=Font(name="Aptos",bold=True,color="FFFFFF",size=10)
        c.fill=PatternFill("solid",fgColor=COR_PRIMARY)
        c.alignment=Alignment(wrap_text=True,vertical="center")
    for row in ws.iter_rows(min_row=2):
        row[1].font=Font(name="Times New Roman",size=13,color="1B374C")
        row[7].font=Font(name="Aptos",size=10,italic=True,color="925A18")
        for c in row:
            c.alignment=Alignment(vertical="center")
        for ci in (0,6):
            row[ci].number_format="#,##0"
    tbl=Table(displayName="ThayerTitulos",ref=f"A1:I{len(entries)+1}")
    tbl.tableStyleInfo=TableStyleInfo(name="TableStyleMedium2",showFirstColumn=False,showLastColumn=False,showRowStripes=True,showColumnStripes=False)
    ws.add_table(tbl)
    return invalid

def fill_duplicates(ws, entries):
    counts=collections.Counter(str(e.get("lemma") or "").strip() for e in entries if e.get("lemma"))
    ws.append(["Título repetido na transcrição","Ocorrências","ID dos registros (amostra)"])
    for i,e in enumerate(ws[1],1):
        e.fill=PatternFill("solid",fgColor=COR_PRIMARY)
        e.font=Font(bold=True,color="FFFFFF")
    groups=collections.defaultdict(list)
    for x in entries:
        k=str(x.get("lemma") or "").strip()
        if k:groups[k].append(str(x.get("id") or ""))
    for lemma,count in sorted(counts.items(),key=lambda x:(-x[1],x[0])):
        if count>1:ws.append([lemma,count,"; ".join(groups[lemma][:15])])
    ws.column_dimensions["A"].width=45
    ws.column_dimensions["B"].width=17
    ws.column_dimensions["C"].width=90
    ws.freeze_panes="A2"
    ws.auto_filter.ref=ws.dimensions

def main():
    req=urllib.request.Request(URL,headers={"User-Agent":"ScripturaLexicon-Thayer-title-index/1.0"})
    with urllib.request.urlopen(req,timeout=120) as r:buf=r.read()
    data=[json.loads(l) for l in buf.decode("utf-8-sig").splitlines() if l.strip()]
    if len(data)<5000:
        raise RuntimeError(f"Fonte digital incompleta: apenas {len(data)} entradas")
    w=Workbook()
    summ=w.active;summ.title="Resumo e método"
    idx=w.create_sheet("Títulos")
    duplicates=w.create_sheet("Títulos repetidos")
    fingerprint=hashlib.sha256(buf).hexdigest()
    fill_summary(summ,data,fingerprint)
    invalid=fill_index(idx,data)
    fill_duplicates(duplicates,data)
    w.properties.title="Thayer — Índice de títulos (inventário da transcrição auxiliar)"
    w.properties.subject="Lista de cabeçalhos e indicação bibliográfica, sem cotejo exaustivo com fac-símile"
    w.properties.creator="Scriptura Lexicon"
    DEST.parent.mkdir(parents=True,exist_ok=True)
    w.save(DEST)
    re=load_workbook(DEST,read_only=True)
    count=re["Títulos"].max_row-1
    if count!=len(data):
        raise RuntimeError(f"Falha na gravação: {count} != {len(data)}")
    print(json.dumps({"saved":str(DEST),"entries":len(data),"distinct_titles":len({str(x.get("lemma") or "").strip() for x in data}),"missing_titles":invalid,"sha256_dataset":fingerprint},ensure_ascii=False))

if __name__=="__main__":main()
