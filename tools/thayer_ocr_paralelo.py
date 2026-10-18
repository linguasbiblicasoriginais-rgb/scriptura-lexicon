#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Extrator de lemas THAYER diretamente de recortes das fotografias, por lotes paralelos.
As leituras do OCR NÃO equivalem à confirmação visual humana.
"""
from __future__ import annotations
from collections import Counter,defaultdict
from datetime import datetime,timezone
import json,os,pathlib,re,subprocess,sys
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
import thayer_indexar_fotografias_v2 as original

ROOT=pathlib.Path(__file__).resolve().parents[1]
CAND=pathlib.Path("/tmp/thayer-candidatos/candidatos.json")
PIECES=pathlib.Path("/tmp/thayer-ocr-results")
OUT=ROOT/"entregas/THAYER_indice_fotografico_paralelo_OCR_provisorio.xlsx"
SUMMARY=ROOT/"entregas/THAYER_indice_fotografico_paralelo_OCR_provisorio.json"
REGEX=re.compile(r"[\u0370-\u03ff\u1f00-\u1fff][\u0370-\u03ff\u1f00-\u1fff\u0300-\u036f᾿᾽'’\-]{0,39}")

def key(r):return f'{r["pagina_pdf"]}:{r["x_scan"]}:{r["y_scan"]}'
def save_json(path,v):
    path.parent.mkdir(parents=True,exist_ok=True)
    path.write_text(json.dumps(v,ensure_ascii=False)+"\n",encoding="utf-8")
def prepare():
    rows,pages=original.extract()
    if len(pages)!=683 or len(rows)<4500:
        raise RuntimeError(f"Inventário geométrico insuficiente: páginas={len(pages)} candidatos={len(rows)}")
    rows.sort(key=lambda v:(v["pagina_pdf"],0 if v["coluna"]=="esquerda" else 1,v["y_scan"],v["x_scan"]))
    save_json(CAND,{"source":"thayersgreekengl0000thay","pages":pages,"candidates":rows})
    print(f"PREPARACAO_OK paginas=683 candidatos={len(rows)}",flush=True)
def tokenize(s):
    s=re.sub(r"^[^Α-Ωα-ω\u0370-\u03ff\u1f00-\u1fff]*","",str(s or ""))
    m=REGEX.search(s)
    return m.group(0).rstrip("'-᾿") if m else ""
def read_crop(pic,r):
    import pytesseract
    from PIL import Image,ImageOps,ImageEnhance
    w,h=pic.size
    kx=w/r["largura_scan"];ky=h/r["altura_scan"]
    x=max(0,int((r["x_scan"]-12)*kx))
    y=max(0,int((r["y_scan"]-12)*ky))
    x2=min(w,int((r["x_scan"]+245)*kx))
    y2=min(h,int((r["y_scan"]+54)*ky))
    if x2<=x or y2<=y:return "", "coordenadas fora da fotografia"
    sample=pic.crop((x,y,x2,y2))
    sample=ImageOps.autocontrast(sample)
    sample=ImageEnhance.Contrast(sample).enhance(1.5)
    # Reconhecimento da PALAVRA em vez de um longo parágrafo reduz o tempo.
    try:
        got=pytesseract.image_to_string(sample,lang="grc",config="--oem 1 --psm 8",timeout=5)
        return tokenize(got),got.replace("\n"," ").strip()[:130]
    except Exception as exc:
        return "",f"OCR não resolvido: {str(exc)[:100]}"
def recognize(part):
    import fitz
    from PIL import Image
    total_parts=8
    v=json.loads(CAND.read_text(encoding="utf-8"))
    groups=defaultdict(list)
    for row in v["candidates"]:
        p=int(row["pagina_pdf"])
        if (p-original.START)%total_parts==part:groups[p].append(row)
    original.download_pdf()
    doc=fitz.open(str(original.PDF))
    out=[];processed=0;found=0
    for p,rows in sorted(groups.items()):
        try:
            pix=doc[p-1].get_pixmap(matrix=fitz.Matrix(4.2,4.2),colorspace=fitz.csGRAY,alpha=False)
            image=Image.frombytes("L",(pix.width,pix.height),pix.samples)
            for row in rows:
                title,excerpt=read_crop(image,row)
                found+=bool(title)
                out.append({"key":key(row),"title":title,"extract":excerpt,"state":
                    "OCR GREGO DA FOTOGRAFIA — NÃO COTEJADO" if title else "LEITURA OCR NÃO OBTIDA"})
        except Exception as error:
            for row in rows:
                out.append({"key":key(row),"title":"","extract":str(error)[:120],"state":"FALHA — REVISAR FOTOGRAFIA"})
        processed+=1
        if processed%14==0:
            print(f"PART-{part}: folhas {processed}/{len(groups)}; OCR={found}/{len(out)}",flush=True)
    doc.close()
    dest=PIECES/f"parte-{part}.json"
    save_json(dest,{"part":part,"pages":processed,"recognized":found,"entries":out})
    print(f"PART-{part}-CONCLUIDA paginas={processed} candidatos={len(out)} grego={found}",flush=True)
def consolidate():
    v=json.loads(CAND.read_text(encoding="utf-8"))
    parts=list(sorted(PIECES.glob("parte-*.json")))
    if len(parts)!=8:raise RuntimeError(f"Apenas {len(parts)} dos 8 lotes foram publicados.")
    greek={};meta={}
    for path in parts:
        chunk=json.loads(path.read_text(encoding="utf-8"))
        for r in chunk["entries"]:
            if r["key"] in greek:raise RuntimeError(f"Duplicação de coordenada: {r['key']}")
            greek[r["key"]]=r["title"]
            meta[r["key"]]=r
    source=v["candidates"]
    if len(greek)!=len(source):raise RuntimeError(f"Linhas OCR {len(greek)} != candidatos {len(source)}")
    for r in source:
        m=meta[key(r)]
        r["titulo_grego_ocr_fotografia"]=m["title"]
        r["status"]=m["state"]
        r["observacoes"]=m["extract"]
    count=sum(bool(r["titulo_grego_ocr_fotografia"]) for r in source)
    by_page=Counter(x["pagina_impressa"] for x in source)
    greek_page=Counter(x["pagina_impressa"] for x in source if x["titulo_grego_ocr_fotografia"])
    w=Workbook();intro=w.active;intro.title="Método e limites"
    notes=[
        ["THAYER — ÍNDICE FOTOGRÁFICO POR OCR PARA CONFERÊNCIA"],
        ["Fonte primária","Fac-símile thayersgreekengl0000thay (edição corrigida do Thayer)"],
        ["Extensão processada","683 páginas lexicográficas impressas / PDF folhas 25–707"],
        ["Quantidade de inícios de artigos CANDIDATOS",len(source)],
        ["Com leitura grega automatizada",count],
        ["Sem leitura grega automatizada",len(source)-count],
        ["Número exato de artigos da edição impressa","NÃO DETERMINADO"],
        ["CONFERÊNCIA VISUAL EXAUSTIVA","NÃO REALIZADA; a lista contém omissões e possíveis falsos positivos"],
        ["Método","OCR com coordenadas da mesma digitalização e segundo OCR grc de palavra isolada da fotografia (8 lotes paralelos)"],
        ["Strong / dataset lexical alternativo","NÃO utilizados para formar nem completar os títulos"],
        ["Observação","Não confundir 5.594 palavras do vocabulário do NT no prefácio com artigos impressos"],
        ["Localização dos títulos","Página impressa, folha do PDF, coluna e coordenadas da fotografia"],
        ["Estado",datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")]]
    for n in notes:intro.append(n)
    ws=w.create_sheet("Títulos OCR — 683 páginas")
    cols=["Nº candidato","Título grego OCR da fotografia","Página impressa","Folha PDF","Coluna",
      "X na imagem","Y na imagem","Trecho do OCR inicial","Leitura OCR do recorte","Estado",
      "Confirmação manual","Abrir fac-símile"]
    ws.append(cols)
    uncertain=w.create_sheet("Pendentes ou ilegíveis");uncertain.append(cols)
    for i,r in enumerate(source,1):
        photo=f"https://archive.org/details/thayersgreekengl0000thay/page/n{r['pagina_pdf']-1}/mode/1up"
        data=[i,r["titulo_grego_ocr_fotografia"],r["pagina_impressa"],r["pagina_pdf"],r["coluna"],
           r["x_scan"],r["y_scan"],r["texto_bruto_inicio"],r["observacoes"],
           r["status"],"PENDENTE","Abrir imagem"]
        ws.append(data)
        c=ws.cell(ws.max_row,12);c.hyperlink=photo;c.font=Font(color="13638C",underline="single")
        if not r["titulo_grego_ocr_fotografia"]:
            uncertain.append(data)
            c=uncertain.cell(uncertain.max_row,12);c.hyperlink=photo;c.font=Font(color="13638C",underline="single")
    coverage=w.create_sheet("Cobertura por página")
    coverage.append(["Página impressa","Folha PDF","Candidatos OCR","Lemas gregos obtidos","Sem leitura","Conferência visual"])
    for p in range(1,684):coverage.append([p,p+24,by_page[p],greek_page[p],by_page[p]-greek_page[p],"PENDENTE"])
    for sh in (ws,uncertain,coverage):
        sh.freeze_panes="C2";sh.auto_filter.ref=sh.dimensions;sh.sheet_view.showGridLines=False
        for c in sh[1]:
            c.font=Font(color="FFFFFF",bold=True)
            c.fill=PatternFill("solid",fgColor="173D54")
            c.alignment=Alignment(wrap_text=True,vertical="center")
        sh.row_dimensions[1].height=32
    for tab in (ws,uncertain):
        for col,size in enumerate([12,42,18,17,14,16,16,65,55,47,24,19],1):
            tab.column_dimensions[get_column_letter(col)].width=size
    for col,size in enumerate([21,17,23,25,18,30],1):
        coverage.column_dimensions[get_column_letter(col)].width=size
    intro.column_dimensions["A"].width=47;intro.column_dimensions["B"].width=118
    for row in intro:
        row[0].font=Font(bold=True,color="173D54")
        if len(row)>1:row[1].alignment=Alignment(wrap_text=True)
    w.properties.title="Thayer — candidatos de cabeçalhos do fac-símile OCR rápido, não conferidos"
    w.properties.creator="Scriptura Lexicon"
    OUT.parent.mkdir(parents=True,exist_ok=True)
    w.save(OUT)
    save_json(SUMMARY,{"source":"facsimile-thayersgreekengl0000thay",
      "pages":683,"candidate_count":len(source),"greek_ocr":count,"unread":len(source)-count,
      "manual_verification":False,"method":"ocr_greek_photographic_crops","entries":source})
    subprocess.run(["git","config","user.name","github-actions[bot]"],check=True,cwd=ROOT)
    subprocess.run(["git","config","user.email","41898282+github-actions[bot]@users.noreply.github.com"],check=True,cwd=ROOT)
    subprocess.run(["git","add",str(OUT.relative_to(ROOT)),str(SUMMARY.relative_to(ROOT))],check=True,cwd=ROOT)
    subprocess.run(["git","commit","-m","THAYER: salvar índice OCR fotográfico de 683 páginas e pendências"],check=True,cwd=ROOT)
    subprocess.run(["git","push","origin","HEAD:chat-gpt-thayer-ocr-paralelo"],check=True,cwd=ROOT)
    print(f"FINAL: {len(source)} candidatos, {count} OCR grego, {len(source)-count} sem leitura",flush=True)
if __name__=="__main__":
    mode=sys.argv[1] if len(sys.argv)>1 else ""
    if mode=="prepare":prepare()
    elif mode=="recognize":recognize(int(sys.argv[2]))
    elif mode=="consolidate":consolidate()
    else:raise SystemExit("Uso: script.py prepare | recognize PART | consolidate")
