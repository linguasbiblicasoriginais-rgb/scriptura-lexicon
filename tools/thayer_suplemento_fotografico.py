#!/usr/bin/env python3
"""Recupera cabeçalhos omitidos por divisão errada de colunas, usando as mesmas fotos.
A saída NÃO afirma revisão visual integral, nem utiliza Strong ou léxicos substitutos.
"""
from __future__ import annotations
import json,urllib.request,subprocess
from concurrent.futures import ProcessPoolExecutor,as_completed
from collections import defaultdict
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import PatternFill,Font,Alignment
import thayer_indexar_fotografias_v2 as base

ROOT=Path(__file__).resolve().parents[1]
OLD_URL="https://raw.githubusercontent.com/linguasbiblicasoriginais-rgb/scriptura-lexicon/chat-gpt-thayer-indice-titulos/entregas/THAYER_inventario_fotografico_candidatos.json"
DEST=ROOT/"entregas/THAYER_cabecalhos_recuperados_coluna_direita.xlsx"
JSON_DEST=ROOT/"entregas/THAYER_cabecalhos_recuperados_coluna_direita.json"
def read_source():
    req=urllib.request.Request(OLD_URL,headers={"User-Agent":"ScripturaLexicon-photographic-columns/1.0"})
    with urllib.request.urlopen(req,timeout=90) as fh:return json.load(fh)
def key(r):return (int(r["pagina_pdf"]),int(r["x_scan"]),int(r["y_scan"]))
def main():
    old=read_source()
    source=old["candidatos"]; seen={key(x) for x in source}
    improved,coverage=base.extract()
    extra=[r for r in improved if key(r) not in seen]
    print(f"ORIGINAL {len(source)} | CORRIGIDO {len(improved)} | SUPLEMENTARES {len(extra)}",flush=True)
    if len(extra)<50:raise RuntimeError("Recuperação insuficiente: confirmar os parâmetros; não publicar.")
    base.download_pdf()
    groups=defaultdict(list)
    for i,r in enumerate(extra):groups[r["pagina_pdf"]].append((i,r))
    passed=0;done=0
    jobs=[(p,[r for _,r in arr]) for p,arr in groups.items()]
    with ProcessPoolExecutor(max_workers=4) as pool:
        mp={pool.submit(base.enrich_single,job):job[0] for job in jobs}
        for fut in as_completed(mp):
            p=mp[fut]
            try:
                _,res=fut.result()
                for (i,r),(head,conf,ocr) in zip(groups[p],res):
                    if head and base.GREK.search(head):
                        extra[i]["titulo_grego_ocr_fotografia"]=head
                        extra[i]["confianca_ocr_grego"]=conf
                        extra[i]["status"]="OCR GREGO DA FOTOGRAFIA — NÃO CONFERIDO MANUALMENTE"
                        extra[i]["observacoes"]=f"Recuperado após corrigir divisão das colunas; texto OCR: {ocr[:110]}"
                        passed+=1
                    else:
                        extra[i]["status"]="SEM LEITURA GREGA CONFIÁVEL — INSPECIONAR FOTO"
            except Exception as err:
                for i,r in groups[p]:
                    extra[i]["status"]="ERRO OCR FOTOGRÁFICO — INSPECIONAR FOTO"
                    extra[i]["observacoes"]=repr(err)[:160]
            done+=1
            if done%35==0:print(f"FOLHAS COMPLEMENTARES {done}/{len(groups)}; OCR legível {passed}/{len(extra)}",flush=True)
    wb=Workbook()
    about=wb.active;about.title="Estado e método"
    for line in [
        ["THAYER — suplemento fotográfico para páginas omitidas"],
        ["Fonte exclusiva para posição e leitura","thayersgreekengl0000thay: fotografias, DjVu XML e OCR grego dos recortes"],
        ["Erro corrigido","A coluna direita de páginas ímpares pode começar antes de 51% da largura total"],
        ["Critério corrigido","Coluna direita detectada a partir de 45% da largura; páginas com zero candidatos receberam reprocessamento"],
        ["Candidatos novos",len(extra)],
        ["Candidatos novos com OCR grego",passed],
        ["Status de conferência","EXTRAÇÃO AUTOMÁTICA; NÃO EQUIVALE A INVENTÁRIO VISUAL CONFERIDO"],
        ["Números Strong utilizados","NÃO"],
        ["Quantitativo definitivo do Thayer","AINDA NÃO DETERMINADO"]
    ]:about.append(line)
    ws=wb.create_sheet("Cabeçalhos recuperados")
    ws.append(["Nº","Leitura grega da fotografia (OCR)","Trecho OCR geométrico","Página impressa","Folha PDF","Coluna","X","Y","Status","Confiança OCR","Link da fotografia"])
    for i,r in enumerate(extra,1):
        url=f'https://archive.org/details/thayersgreekengl0000thay/page/n{r["pagina_pdf"]-1}/mode/1up'
        ws.append([i,r["titulo_grego_ocr_fotografia"],r["texto_bruto_inicio"],r["pagina_impressa"],r["pagina_pdf"],r["coluna"],r["x_scan"],r["y_scan"],r["status"],r["confianca_ocr_grego"],"Ver imagem"])
        c=ws.cell(i+1,11);c.hyperlink=url;c.font=Font(color="14638D",underline="single")
    for tab in [ws,about]:
        tab.sheet_view.showGridLines=False
        for c in tab[1]:
            c.font=Font(bold=True,color="FFFFFF")
            c.fill=PatternFill("solid",fgColor="17384F")
    ws.freeze_panes="C2";ws.auto_filter.ref=ws.dimensions
    from openpyxl.utils import get_column_letter
    for j,w in enumerate([12,37,57,20,18,16,12,12,50,20,22],1):ws.column_dimensions[get_column_letter(j)].width=w
    about.column_dimensions["A"].width=38;about.column_dimensions["B"].width=110
    DEST.parent.mkdir(parents=True,exist_ok=True)
    wb.save(DEST)
    JSON_DEST.write_text(json.dumps({"source_photographic":"thayersgreekengl0000thay","old_candidates":len(source),
      "improved_candidates":len(improved),"delta_candidates":len(extra),"greek_ocr_obtained":passed,
      "state":"automatic_not_manually_checked","entries":extra},ensure_ascii=False)+"\n",encoding="utf-8")
    subprocess.run(["git","config","user.name","github-actions[bot]"],check=True,cwd=ROOT)
    subprocess.run(["git","config","user.email","41898282+github-actions[bot]@users.noreply.github.com"],check=True,cwd=ROOT)
    subprocess.run(["git","add",str(DEST.relative_to(ROOT)),str(JSON_DEST.relative_to(ROOT))],check=True,cwd=ROOT)
    subprocess.run(["git","commit","-m","THAYER: recuperar candidatos omitidos no fac-símile, com OCR grego"],check=True,cwd=ROOT)
    subprocess.run(["git","push","origin","HEAD:chat-gpt-thayer-facsimile-apoio"],check=True,cwd=ROOT)
    print("FINAL RECUPERADOS",len(extra),"LEITURAS_OCR",passed,flush=True)
if __name__=="__main__":main()
